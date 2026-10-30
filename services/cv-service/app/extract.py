import io
import re

MAX_TEXT_CHARS = 20_000
MAX_PDF_PAGES = 20


class UnsupportedFormat(ValueError):
    """The file is not a readable PDF, DOCX or UTF-8 text file."""


class EmptyDocument(ValueError):
    """The file was readable but contains no text (e.g. a scanned PDF)."""


def detect_format(data: bytes) -> str:
    """Detect the format from the first bytes, never from the file name."""
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "docx"
    return "txt"


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedFormat("encrypted PDFs are not supported")
        return "\n".join(p.extract_text() or "" for p in reader.pages[:MAX_PDF_PAGES])
    except UnsupportedFormat:
        raise
    except Exception as e:  # untrusted input: any parser failure means "unreadable"
        raise UnsupportedFormat("could not read the PDF") from e


def _read_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as e:
        raise UnsupportedFormat("could not read the DOCX") from e


def _read_txt(data: bytes) -> str:
    if b"\x00" in data:
        raise UnsupportedFormat("binary file")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise UnsupportedFormat("text files must be UTF-8") from e


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_URL = re.compile(r"https?://\S+|www\.\S+")
_PHONE = re.compile(r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)")


def _redact_phone(match: re.Match) -> str:
    # Require 9+ digits so that date ranges like "2019 - 2022" are kept
    if sum(c.isdigit() for c in match.group()) >= 9:
        return "[phone]"
    return match.group()


def redact_pii(text: str) -> str:
    """Replace emails, URLs and phone numbers. Names and addresses are NOT detected."""
    text = _EMAIL.sub("[email]", text)
    text = _URL.sub("[url]", text)
    return _PHONE.sub(_redact_phone, text)


def process_cv(data: bytes) -> tuple[str, str]:
    """Return (format, redacted text). Raises UnsupportedFormat or EmptyDocument."""
    fmt = detect_format(data)
    raw = {"pdf": _read_pdf, "docx": _read_docx, "txt": _read_txt}[fmt](data)
    text = redact_pii(_clean(raw))[:MAX_TEXT_CHARS].strip()
    if not text:
        raise EmptyDocument("no text found; scanned or image-only files are not supported")
    return fmt, text
